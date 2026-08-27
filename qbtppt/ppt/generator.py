import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from win32com.client import Dispatch

from qbtppt.models import AppSettings, BonusData, SelectionOptions, TossupData
from qbtppt.text.processing import after_substring, ensure_sentence_end, split_tossup
from qbtppt.trivia.categories import category_value
from qbtppt.trivia.difficulties import DIFFICULTY_LABELS
from qbtppt.trivia.distributions import NO_DISTRIBUTION, PACKET_SOURCE_MODE


def format_categories(categories: list[object]) -> str:
    return ", ".join([category_value(category) for category in categories]) if categories else "Any"


def rgb_from_hex(hex_color: str) -> RGBColor:
    cleaned = hex_color.lstrip("#")
    return RGBColor(
        int(cleaned[0:2], 16),
        int(cleaned[2:4], 16),
        int(cleaned[4:6], 16),
    )


def apply_content_slide_style(slide, settings: AppSettings) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb_from_hex(settings.question_background_color)


def add_question_textbox(ppt: Presentation, slide, settings: AppSettings):
    slide_width = ppt.slide_width
    slide_height = ppt.slide_height
    left = Inches(settings.question_margin_left)
    top = Inches(settings.question_margin_top)
    width = slide_width - left - Inches(settings.question_margin_right)
    height = slide_height - top - Inches(settings.question_margin_bottom)

    textbox = slide.shapes.add_textbox(left, top, width, height)
    textbox.text_frame.margin_left = Inches(settings.question_internal_margin)
    textbox.text_frame.margin_right = Inches(settings.question_internal_margin)
    textbox.text_frame.margin_top = Inches(settings.question_internal_margin)
    textbox.text_frame.margin_bottom = Inches(settings.question_internal_margin)
    return textbox


def style_question_paragraph(paragraph, settings: AppSettings) -> None:
    paragraph.space_after = Pt(settings.question_paragraph_spacing_after)
    paragraph.line_spacing = settings.question_line_spacing
    for run in paragraph.runs:
        run.font.name = settings.question_font_family
        run.font.size = Pt(settings.question_max_font_size)
        run.font.color.rgb = rgb_from_hex(settings.question_text_color)


def add_styled_question_slide(ppt: Presentation, settings: AppSettings, shape_name: str, lines: list[str], notes: str) -> None:
    slide = ppt.slides.add_slide(ppt.slide_layouts[6])
    apply_content_slide_style(slide, settings)
    slide.notes_slide.notes_text_frame.text = notes

    textbox = add_question_textbox(ppt, slide, settings)
    textbox.name = shape_name
    text_frame = textbox.text_frame
    text_frame.clear()

    for index, line in enumerate(lines):
        paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        paragraph.text = line
        style_question_paragraph(paragraph, settings)
    text_frame.fit_text(max_size=settings.question_max_font_size, font_family=settings.question_font_family)


def add_bonus_slides(ppt: Presentation, settings: AppSettings, bonus: BonusData, bonus_number: int) -> None:
    bonus_lines = [f"Bonus {bonus_number}", "", bonus.leadin_sanitized]
    for part_index, part in enumerate(bonus.parts_sanitized, start=1):
        value = bonus.values[part_index - 1] if bonus.values and part_index <= len(bonus.values) else 10
        bonus_lines.append(f"[{value}] {ensure_sentence_end(part)}")
        if part_index <= len(bonus.answers_sanitized):
            bonus_lines.append(f"Answer: {bonus.answers_sanitized[part_index - 1]}")

    answer_lines = [
        f"{index + 1}. {answer}"
        for index, answer in enumerate(bonus.answers_sanitized)
    ]
    add_styled_question_slide(
        ppt,
        settings,
        f"RevealText_Bonus_{bonus_number}",
        bonus_lines,
        "Answers:\n" + "\n".join(answer_lines),
    )


def style_textbox_runs(textbox, settings: AppSettings, max_font_size: int) -> None:
    for paragraph in textbox.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.name = settings.question_font_family
            run.font.size = Pt(max_font_size)
            run.font.color.rgb = rgb_from_hex(settings.question_text_color)


def add_answer_slide(ppt: Presentation, settings: AppSettings, answer: str, subtitle: str = "") -> None:
    slide = ppt.slides.add_slide(ppt.slide_layouts[6])
    apply_content_slide_style(slide, settings)

    answer_height = Inches(3.15 if subtitle else 6.85)
    answer_box = slide.shapes.add_textbox(Inches(0.55), Inches(0.35), Inches(8.9), answer_height)
    answer_box.text_frame.margin_left = Inches(0.05)
    answer_box.text_frame.margin_right = Inches(0.05)
    answer_box.text_frame.margin_top = Inches(0)
    answer_box.text_frame.margin_bottom = Inches(0)
    answer_box.text_frame.clear()
    answer_paragraph = answer_box.text_frame.paragraphs[0]
    answer_paragraph.text = answer
    style_textbox_runs(answer_box, settings, 34)
    answer_box.text_frame.fit_text(max_size=34, font_family=settings.question_font_family)

    if not subtitle:
        return

    subtitle_box = slide.shapes.add_textbox(Inches(0.7), Inches(3.6), Inches(8.6), Inches(3.75))
    subtitle_box.text_frame.margin_left = Inches(0.05)
    subtitle_box.text_frame.margin_right = Inches(0.05)
    subtitle_box.text_frame.margin_top = Inches(0)
    subtitle_box.text_frame.margin_bottom = Inches(0)
    subtitle_box.text_frame.clear()
    subtitle_paragraph = subtitle_box.text_frame.paragraphs[0]
    subtitle_paragraph.text = subtitle
    style_textbox_runs(subtitle_box, settings, 22)
    subtitle_box.text_frame.fit_text(max_size=22, font_family=settings.question_font_family)


def config_slide_text(options: SelectionOptions, tossups: list[TossupData], bonuses: list[BonusData]) -> str:
    if options.source_mode == PACKET_SOURCE_MODE:
        return (
            f"Source:\n"
            f"{options.source_mode}\n\n"
            f"Set:\n"
            f"{options.set_name or 'Unknown'}\n\n"
            f"Packet:\n"
            f"{options.packet_number if options.packet_number is not None else 'Unknown'}\n\n"
            f"Tossups:\n"
            f"{len(tossups)}\n\n"
            f"Bonuses:\n"
            f"{len(bonuses)}"
        )

    return (
        f"Source:\n"
        f"{options.source_mode}\n\n"
        f"Difficulties:\n"
        f"{', '.join([DIFFICULTY_LABELS[diff] for diff in options.difficulties]) if options.difficulties else 'Any'}\n\n"
        f"Years:\n"
        f"{options.min_year}-{options.max_year}\n\n"
        f"Distribution:\n"
        f"{options.distribution}\n\n"
        f"Categories:\n"
        f"{format_categories(options.categories) if options.distribution == NO_DISTRIBUTION else options.distribution}\n\n"
        f"Bonuses:\n"
        f"{'Yes' if options.include_bonuses else 'No'}"
    )


def add_reveal_animations(pptx_path: str, settings: AppSettings) -> None:
    powerpoint = Dispatch("PowerPoint.Application")
    presentation = powerpoint.Presentations.Open(os.path.abspath(pptx_path), WithWindow=False)

    try:
        for slide in presentation.Slides:
            for shape in slide.Shapes:
                if not shape.Name.startswith("RevealText"):
                    continue
                animation_settings = shape.AnimationSettings
                animation_settings.Animate = True
                animation_settings.TextLevelEffect = settings.animate_by_first_level
                animation_settings.TextUnitEffect = settings.animate_by_word
                animation_settings.EntryEffect = settings.effect_appear
        presentation.Save()
    finally:
        presentation.Close()
        powerpoint.Quit()


def build_powerpoint(
    tossups: list[TossupData],
    bonuses: list[BonusData],
    options: SelectionOptions,
    settings: AppSettings,
    output_path: str | None = None,
) -> str:
    deck_path = output_path or settings.default_output_path
    questions = [split_tossup(tossup.question_sanitized) for tossup in tossups]
    answers = [tossup.answer_sanitized for tossup in tossups]

    ppt = Presentation()
    slide = ppt.slides.add_slide(ppt.slide_layouts[0])
    slide.shapes.title.text = settings.deck_title
    slide.placeholders[1].text = settings.deck_subtitle

    slide = ppt.slides.add_slide(ppt.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(0.35), Inches(0.05), Inches(9.3), Inches(7.6))
    text_frame = textbox.text_frame
    text_frame.margin_left = Inches(0.05)
    text_frame.margin_right = Inches(0.05)
    text_frame.margin_top = Inches(0)
    text_frame.margin_bottom = Inches(0)
    text_frame.clear()
    paragraph = text_frame.paragraphs[0]
    paragraph.text = config_slide_text(options, tossups, bonuses)
    for run in paragraph.runs:
        run.font.name = settings.question_font_family
        run.font.size = Pt(15)
        run.font.color.rgb = rgb_from_hex(settings.question_text_color)
    text_frame.fit_text(max_size=15, font_family=settings.question_font_family)

    for q, tossup_parts in enumerate(questions):
        note = ""
        subtitle = ""
        if "[" in answers[q] and "]" in answers[q]:
            subtitle = answers[q][answers[q].find("["):answers[q].find("]") + 1]
            answers[q] = answers[q][:answers[q].find("[")]
        if tossup_parts and tossup_parts[0].startswith("[Note"):
            note_end = tossup_parts[0].find("]")
            note = tossup_parts[0][:note_end + 1]
            tossup_parts[0] = after_substring(tossup_parts[0], "]")

        if note:
            notes = note + f"\n Answer: {answers[q]}\n{subtitle}"
        else:
            notes = f"Answer: {answers[q]}\n{subtitle}"
        add_styled_question_slide(
            ppt,
            settings,
            f"RevealText_Tossup_{q + 1}",
            [ensure_sentence_end(tossup_part) for tossup_part in tossup_parts],
            notes,
        )

        add_answer_slide(ppt, settings, answers[q], subtitle)

        if q < len(bonuses):
            add_bonus_slides(ppt, settings, bonuses[q], q + 1)

    for bonus_index, bonus in enumerate(bonuses[len(questions):], start=len(questions) + 1):
        add_bonus_slides(ppt, settings, bonus, bonus_index)

    ppt.save(deck_path)
    return deck_path
